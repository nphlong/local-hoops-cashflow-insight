from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

VND_RE = re.compile(r"[^\d.-]")
ACCOUNT_RE = re.compile(r"Account No:\s*([0-9]+)", re.I)
FROM_RE = re.compile(r"From:(\d{2}/\d{2}/\d{4})")
TO_RE = re.compile(r"To:(\d{2}/\d{2}/\d{4})")

HEADER_ALIASES = {
    "stt": "seq",
    "ngay giao dich": "booked",
    "ngay hach toan": "posted",
    "phat sinh no": "debit",
    "phat sinh co": "credit",
    "so du": "balance",
    "don vi thu huong/don vi chuyen": "counterparty",
    "don vi thu huong": "counterparty",
    "noi dung": "description",
    "but toan": "bank_ref",
}


def _fold(text: str) -> str:
    swapped = (text or "").replace("đ", "d").replace("Đ", "D")
    normalized = unicodedata.normalize("NFD", swapped)
    stripped = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", stripped).strip().lower()


def parse_amount(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    return float(VND_RE.sub("", text) or 0)


def parse_datetime(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def iso_date(value: datetime | None) -> str | None:
    return value.strftime("%Y-%m-%d") if value else None


def iso_datetime(value: datetime | None) -> str | None:
    return value.strftime("%Y-%m-%dT%H:%M:%S") if value else None


@dataclass
class BronzeRow:
    seq: int
    booked_at: str | None
    posted_on: str | None
    debit: float
    credit: float
    balance_after: float | None
    counterparty: str
    description: str
    bank_ref: str


@dataclass
class BronzeFile:
    filename: str
    account_no: str | None = None
    currency: str = "VND"
    statement_from: str | None = None
    statement_to: str | None = None
    closing_balance: float | None = None
    credit_total: float | None = None
    rows: list[BronzeRow] = field(default_factory=list)


def _cell(sheet, row: int, col: str):
    return sheet[f"{col}{row}"].value


def _row_blob(sheet, row_idx: int, max_col: int) -> str:
    parts = []
    for col_idx in range(1, max_col + 1):
        value = sheet.cell(row_idx, col_idx).value
        if value not in (None, ""):
            parts.append(str(value))
    return " ".join(parts)


def _header_map(sheet, row_idx: int, max_col: int) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for col_idx in range(1, max_col + 1):
        raw = sheet.cell(row_idx, col_idx).value
        if raw in (None, ""):
            continue
        key = _fold(str(raw))
        field = HEADER_ALIASES.get(key)
        if field:
            mapping[field] = get_column_letter(col_idx)
    return mapping


def parse_workbook(path: Path) -> BronzeFile:
    wb = load_workbook(path, data_only=True, read_only=False)
    sheet = wb[wb.sheetnames[0]]
    statement = BronzeFile(filename=path.name)
    header_row = None
    columns: dict[str, str] = {}
    max_row = sheet.max_row or 1
    max_col = sheet.max_column or 20

    for row_idx in range(1, max_row + 1):
        blob = _row_blob(sheet, row_idx, max_col)
        account = ACCOUNT_RE.search(blob)
        if account:
            statement.account_no = account.group(1)
        frm = FROM_RE.search(blob)
        too = TO_RE.search(blob)
        if frm:
            statement.statement_from = iso_date(parse_datetime(frm.group(1)))
        if too:
            statement.statement_to = iso_date(parse_datetime(too.group(1)))
        mapping = _header_map(sheet, row_idx, max_col)
        if mapping.get("seq") and mapping.get("booked"):
            header_row = row_idx
            columns = mapping
            break

    if header_row is None or "credit" not in columns:
        wb.close()
        raise ValueError(f"No transaction header row in {path.name}")

    def value(row_idx: int, field: str):
        col = columns.get(field)
        return _cell(sheet, row_idx, col) if col else None

    extracted: list[BronzeRow] = []
    for row_idx in range(header_row + 1, max_row + 1):
        seq_raw = value(row_idx, "seq")
        seq_text = str(seq_raw).strip() if seq_raw is not None else ""
        folded = _fold(seq_text)
        if folded.startswith("so du cuoi"):
            closing = parse_amount(seq_text.split(":")[-1] if ":" in seq_text else value(row_idx, "balance"))
            if closing:
                statement.closing_balance = closing
            continue
        if not seq_text.isdigit():
            credit = parse_amount(value(row_idx, "credit"))
            if credit and extracted and statement.credit_total is None:
                statement.credit_total = credit
            continue
        booked = parse_datetime(value(row_idx, "booked"))
        posted = parse_datetime(value(row_idx, "posted")) or booked
        extracted.append(
            BronzeRow(
                seq=int(float(seq_text)),
                booked_at=iso_datetime(booked),
                posted_on=iso_date(posted),
                debit=parse_amount(value(row_idx, "debit")),
                credit=parse_amount(value(row_idx, "credit")),
                balance_after=parse_amount(value(row_idx, "balance")) or None,
                counterparty=str(value(row_idx, "counterparty") or "").strip(),
                description=str(value(row_idx, "description") or "").strip(),
                bank_ref=str(value(row_idx, "bank_ref") or "").strip(),
            )
        )

    wb.close()
    extracted.reverse()
    statement.rows = extracted
    if statement.credit_total is None:
        statement.credit_total = sum(row.credit for row in extracted)
    if statement.closing_balance is None and extracted:
        statement.closing_balance = extracted[-1].balance_after
    return statement
